"""
对比 nsys_logs_dy (动态图) 和 nsys_logs_dy2stcinn (CINN融合) 的 GPU kernel 性能,
找出 CINN 融合收益最大的 case, 量化融合提速效果.

思路:
  1. 从两组 log 的 cuda_gpu_kern_sum 段提取所有 GPU kernel 耗时
  2. 将 kernel 分为 "计算密集型(cublas/cudnn)" 和 "访存密集型(其余)"
     - 计算密集型 kernel 两种模式都走相同 library, 不参与 CINN 融合
     - 访存密集型 kernel = Phi逐算子 kernel (dy) vs CINN融合 kernel (cinn)
  3. 对比: dy 访存密集 kernel 总耗时 vs cinn 访存密集 kernel 总耗时
     speedup = (dy_membnd - cinn_membnd) / dy_membnd * 100%

用法:
    python compare_dy_vs_cinn.py

输出:
    dy_vs_cinn_comparison.xlsx
"""

import os
import re
import glob
import csv
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ===== 配置 =====
DY_LOG_DIR = "nsys_logs_dy"
CINN_LOG_DIR = "nsys_logs_dy2stcinn"
OUTPUT_FILE = "dy_vs_cinn_comparison.xlsx"
# ================

# ----------------------------------------------------------------
# 计算密集型 kernel 的匹配模式 (cublas/cudnn, 两种模式都一样, 不可融合)
# ----------------------------------------------------------------
COMPUTE_BOUND_PATTERNS = [
    re.compile(r"sgemm"),
    re.compile(r"hgemm"),
    re.compile(r"dgemm"),
    re.compile(r"igemm"),
    re.compile(r"splitKreduce"),
    re.compile(r"cublasLt::"),
    re.compile(r"cublas"),
    re.compile(r"cudnn::"),
    re.compile(r"cutlass"),
    re.compile(r"volta_"),
    re.compile(r"ampere_s?gemm"),
    re.compile(r"turing_"),
    re.compile(r"sm\d+_xmma"),
    re.compile(r"implicit_convolve"),
    re.compile(r"winograd"),
    re.compile(r"fft2d"),
    re.compile(r"_convolution"),
]


def is_compute_bound(kernel_name):
    for pat in COMPUTE_BOUND_PATTERNS:
        if pat.search(kernel_name):
            return True
    return False


def is_cinn_kernel(kernel_name):
    return kernel_name.startswith("fn_")


# ----------------------------------------------------------------
# CINN kernel 融合类别分析
# ----------------------------------------------------------------
SHORT_TO_FULL = {
    "full": "fill_constant", "sum": "reduce_sum", "r_max": "reduce_max",
    "r_min": "reduce_min", "prod": "reduce_prod", "add": "elementwise_add",
    "mul": "elementwise_mul", "sub": "subtract", "div": "divide",
    "bc": "broadcast_to", "gs": "generate_shape", "yield": "yield_store",
}

OP_CATEGORY = {}
for op in [
    "exp", "erf", "sqrt", "log", "floor", "ceil", "rint", "round", "tanh",
    "log2", "log10", "trunc", "cos", "cosh", "tan", "sin", "sinh",
    "acos", "acosh", "asin", "asinh", "atan", "atanh",
    "bitwise_not", "negative", "identity", "sign", "abs", "rsqrt",
    "sigmoid", "cbrt", "clz", "popc", "isnan", "isfinite", "isinf",
    "scale", "const_scalar", "fill_constant", "assign_value",
    "squeeze", "expand_dims", "reshape", "cast", "yield_store",
    "arange", "logical_not", "tril", "assign_out_", "isclose",
    "select", "slice_assign", "reciprocal", "logical_right_shift",
    "generate_shape",
    "elementwise_add", "atan2", "elementwise_mul", "subtract", "divide",
    "floor_divide", "mod", "remainder", "max", "min", "pow",
    "logical_and", "logical_or", "logical_xor",
    "greater_than", "less_than", "equal", "not_equal",
    "greater_equal", "less_equal",
    "bitwise_or", "bitwise_xor", "bitwise_and",
    "left_shift", "right_shift",
]:
    OP_CATEGORY[op] = "Elementwise"
for op in ["broadcast_to"]:
    OP_CATEGORY[op] = "Broadcast"
for op in ["concat", "reverse", "transpose", "slice", "gather",
           "scatter_assign", "bitcast_convert", "lookup_table", "one_hot", "gather_nd"]:
    OP_CATEGORY[op] = "Injective"
for op in ["reduce_sum", "reduce_prod", "variance", "argmax", "argmin",
           "reduce_max", "reduce_min", "reduce_all", "reduce_any"]:
    OP_CATEGORY[op] = "Reduce"

TOKEN_CATEGORY = dict(OP_CATEGORY)
for short, full in SHORT_TO_FULL.items():
    if full in OP_CATEGORY:
        TOKEN_CATEGORY[short] = OP_CATEGORY[full]

ALL_TOKENS = sorted(TOKEN_CATEGORY.keys(), key=len, reverse=True)


def parse_cinn_kernel_categories(kernel_name):
    clean = re.sub(r"_kernel$", "", kernel_name)
    m = re.match(r"fn_(.*?)__\d+", clean)
    if not m:
        return set()
    ops_str = m.group(1)
    categories = set()
    remaining = ops_str
    while remaining:
        if remaining.startswith("_"):
            remaining = remaining[1:]
            continue
        matched = False
        for token in ALL_TOKENS:
            if remaining.startswith(token):
                rest = remaining[len(token):]
                if rest == "" or rest.startswith("_"):
                    if token in TOKEN_CATEGORY:
                        categories.add(TOKEN_CATEGORY[token])
                    remaining = rest
                    matched = True
                    break
        if not matched:
            idx = remaining.find("_")
            if idx == -1:
                break
            remaining = remaining[idx:]
    return categories


# ----------------------------------------------------------------
# Log 解析
# ----------------------------------------------------------------
def extract_gpu_kernels(log_path):
    """从 log 的 cuda_gpu_kern_sum 段提取 GPU kernel 信息"""
    kernels = []
    in_section = False

    with open(log_path, "r", errors="replace") as f:
        for line in f:
            s = line.strip()

            if "cuda_gpu_kern_sum" in s:
                in_section = True
                continue

            if in_section and s.startswith("Time (%),"):
                continue

            if in_section and (s == "" or s.startswith("Processing")):
                if kernels:
                    break
                in_section = False
                continue

            if not in_section:
                continue

            if not s or not s[0].isdigit():
                continue
            try:
                reader = csv.reader(io.StringIO(s))
                fields = next(reader)
                if len(fields) < 9:
                    continue
                kernels.append({
                    "name": fields[-1].strip(),
                    "total_ns": int(fields[1]),
                    "instances": int(fields[2]),
                    "avg_ns": float(fields[3]),
                    "time_pct": float(fields[0]),
                })
            except (ValueError, IndexError, StopIteration):
                continue

    # 兜底: 直接搜索 Instances 表头
    if not kernels:
        with open(log_path, "r", errors="replace") as f:
            found = False
            for line in f:
                s = line.strip()
                if "Time (%),Total Time (ns),Instances" in s:
                    found = True
                    continue
                if found and s and s[0].isdigit():
                    try:
                        reader = csv.reader(io.StringIO(s))
                        fields = next(reader)
                        if len(fields) >= 9:
                            kernels.append({
                                "name": fields[-1].strip(),
                                "total_ns": int(fields[1]),
                                "instances": int(fields[2]),
                                "avg_ns": float(fields[3]),
                                "time_pct": float(fields[0]),
                            })
                    except (ValueError, IndexError, StopIteration):
                        continue
                elif found and (s == "" or s.startswith("Processing")):
                    if kernels:
                        break
                    found = False
    return kernels


def analyze_kernels(kernels):
    total_ns = compute_ns = membnd_ns = cinn_ns = 0
    total_inst = membnd_inst = cinn_inst = 0
    cinn_categories = set()

    for k in kernels:
        total_ns += k["total_ns"]
        total_inst += k["instances"]
        if is_compute_bound(k["name"]):
            compute_ns += k["total_ns"]
        elif is_cinn_kernel(k["name"]):
            cinn_ns += k["total_ns"]
            cinn_inst += k["instances"]
            membnd_ns += k["total_ns"]
            membnd_inst += k["instances"]
            cinn_categories |= parse_cinn_kernel_categories(k["name"])
        else:
            membnd_ns += k["total_ns"]
            membnd_inst += k["instances"]

    return {
        "total_ns": total_ns, "compute_ns": compute_ns,
        "membnd_ns": membnd_ns, "cinn_ns": cinn_ns,
        "total_inst": total_inst, "membnd_inst": membnd_inst,
        "cinn_inst": cinn_inst, "cinn_categories": cinn_categories,
    }


def main():
    dy_logs = {os.path.splitext(os.path.basename(f))[0]: f
               for f in glob.glob(os.path.join(DY_LOG_DIR, "*.log"))}
    cinn_logs = {os.path.splitext(os.path.basename(f))[0]: f
                 for f in glob.glob(os.path.join(CINN_LOG_DIR, "*.log"))}

    common = sorted(set(dy_logs.keys()) & set(cinn_logs.keys()))
    print(f"[INFO] dy: {len(dy_logs)}, cinn: {len(cinn_logs)}, 匹配: {len(common)}")

    if not common:
        print("[ERROR] 无匹配 case")
        return

    results = []
    all_cats = set()

    for idx, case in enumerate(common, 1):
        if idx % 100 == 0:
            print(f"  [{idx}/{len(common)}]...")

        dy_k = extract_gpu_kernels(dy_logs[case])
        ci_k = extract_gpu_kernels(cinn_logs[case])
        if not dy_k and not ci_k:
            continue

        dy_s = analyze_kernels(dy_k)
        ci_s = analyze_kernels(ci_k)

        dy_mb = dy_s["membnd_ns"]
        ci_mb = ci_s["membnd_ns"]

        if dy_mb > 0 and ci_mb > 0:
            sp_pct = (dy_mb - ci_mb) / dy_mb * 100.0
            sp_ratio = dy_mb / ci_mb
        elif dy_mb > 0:
            sp_pct, sp_ratio = 100.0, float("inf")
        else:
            sp_pct, sp_ratio = 0.0, 1.0

        all_cats |= ci_s["cinn_categories"]

        results.append({
            "case": case,
            "dy_total_ns": dy_s["total_ns"], "dy_compute_ns": dy_s["compute_ns"],
            "dy_membnd_ns": dy_mb, "dy_total_inst": dy_s["total_inst"],
            "dy_membnd_inst": dy_s["membnd_inst"],
            "ci_total_ns": ci_s["total_ns"], "ci_compute_ns": ci_s["compute_ns"],
            "ci_membnd_ns": ci_mb, "ci_fn_ns": ci_s["cinn_ns"],
            "ci_total_inst": ci_s["total_inst"], "ci_membnd_inst": ci_s["membnd_inst"],
            "ci_fn_inst": ci_s["cinn_inst"],
            "sp_pct": sp_pct, "sp_ratio": sp_ratio,
            "cats": ci_s["cinn_categories"],
            "has_cinn": ci_s["cinn_ns"] > 0,
            "dy_kernels": dy_k, "ci_kernels": ci_k,
        })

    with_cinn = [r for r in results if r["has_cinn"]]
    print(f"\n[INFO] 有 CINN 融合 kernel: {len(with_cinn)} / {len(results)}")

    if with_cinn:
        sps = [r["sp_pct"] for r in with_cinn]
        print(f"[RESULT] 访存密集型 kernel 加速:")
        print(f"  平均: {sum(sps)/len(sps):.1f}%")
        print(f"  中位数: {sorted(sps)[len(sps)//2]:.1f}%")
        print(f"  >0%: {sum(1 for s in sps if s>0)}/{len(sps)}")
        print(f"  >=50%: {sum(1 for s in sps if s>=50)}/{len(sps)}")
        print(f"  类别: {sorted(all_cats)}")

    # ===== Excel =====
    wb = Workbook()
    hf = Font(bold=True, color="FFFFFF")
    hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    gfill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    bfill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    bd = Border(left=Side("thin"), right=Side("thin"), top=Side("thin"), bottom=Side("thin"))

    def wh(ws, headers):
        for c, h in enumerate(headers, 1):
            cell = ws.cell(1, c, h)
            cell.font = hf; cell.fill = hfill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = bd

    # Sheet 1: Per-Case
    ws1 = wb.active
    ws1.title = "Per-Case Comparison"
    h1 = [
        "Case",
        "dy 总耗时(ns)", "dy 计算密集(ns)", "dy 访存密集(ns)",
        "dy 总launch", "dy 访存launch",
        "cinn 总耗时(ns)", "cinn 计算密集(ns)", "cinn 访存密集(ns)",
        "cinn fn_耗时(ns)", "cinn 总launch", "cinn 访存launch", "cinn fn_launch",
        "访存加速(%)", "访存加速倍数",
        "E", "B", "I", "R", "融合类别",
    ]
    wh(ws1, h1)

    rsorted = sorted(results, key=lambda r: r["sp_pct"], reverse=True)
    for i, r in enumerate(rsorted, 2):
        cats = r["cats"]
        row = [
            r["case"].replace("__", "/"),
            r["dy_total_ns"], r["dy_compute_ns"], r["dy_membnd_ns"],
            r["dy_total_inst"], r["dy_membnd_inst"],
            r["ci_total_ns"], r["ci_compute_ns"], r["ci_membnd_ns"],
            r["ci_fn_ns"], r["ci_total_inst"], r["ci_membnd_inst"], r["ci_fn_inst"],
            round(r["sp_pct"], 1),
            round(r["sp_ratio"], 2) if r["sp_ratio"] != float("inf") else "INF",
            "Y" if "Elementwise" in cats else "",
            "Y" if "Broadcast" in cats else "",
            "Y" if "Injective" in cats else "",
            "Y" if "Reduce" in cats else "",
            "+".join(sorted(cats)) if cats else "",
        ]
        for c, v in enumerate(row, 1):
            cell = ws1.cell(i, c, v)
            cell.border = bd
            if c == 14 and isinstance(v, (int, float)):
                cell.fill = gfill if v >= 50 else (bfill if v < 0 else PatternFill())
            if c in (16, 17, 18, 19) and v == "Y":
                cell.fill = gfill
                cell.alignment = Alignment(horizontal="center")

    # Sheet 2: Summary
    ws2 = wb.create_sheet("Summary")
    sdata = [["指标", "全部 case", "有 CINN kernel 的 case"]]
    sdata.append(["Case 数量", len(results), len(with_cinn)])
    if with_cinn:
        csps = [r["sp_pct"] for r in with_cinn]
        asps = [r["sp_pct"] for r in results if r["dy_membnd_ns"] > 0]
        sdata.append(["平均加速(%)", f"{sum(asps)/len(asps):.1f}" if asps else "N/A",
                       f"{sum(csps)/len(csps):.1f}"])
        sdata.append(["中位数加速(%)",
                       f"{sorted(asps)[len(asps)//2]:.1f}" if asps else "N/A",
                       f"{sorted(csps)[len(csps)//2]:.1f}"])
        sdata.append(["加速>0%", sum(1 for s in asps if s > 0), sum(1 for s in csps if s > 0)])
        sdata.append(["加速>=50%", sum(1 for s in asps if s >= 50), sum(1 for s in csps if s >= 50)])
    sdata.append([])
    sdata.append(["CINN 融合涵盖算子类别", ", ".join(sorted(all_cats))])
    sdata.append(["类别数量", len(all_cats)])
    sdata.append([])
    sdata.append(["融合类别组合", "Case 数", "平均加速(%)"])
    combos = {}
    for r in with_cinn:
        k = "+".join(sorted(r["cats"])) if r["cats"] else "None"
        combos.setdefault(k, []).append(r["sp_pct"])
    for k in sorted(combos):
        v = combos[k]
        sdata.append([k, len(v), f"{sum(v)/len(v):.1f}"])

    for ri, row in enumerate(sdata, 1):
        for ci, v in enumerate(row, 1):
            cell = ws2.cell(ri, ci, v)
            cell.border = bd
            if ri == 1:
                cell.font = hf; cell.fill = hfill

    # Sheet 3: Top 30 Best (含 kernel 明细)
    ws3 = wb.create_sheet("Top30 Best Fusion")
    wh(ws3, ["Case", "访存加速(%)", "dy 访存 kernel 明细", "cinn 访存 kernel 明细"])

    top30 = [r for r in rsorted if r["has_cinn"]][:30]
    for i, r in enumerate(top30, 2):
        dy_str = "\n".join(
            f"{k['name'][:100]}: {k['total_ns']}ns x{k['instances']}"
            for k in r["dy_kernels"] if not is_compute_bound(k["name"])
        )
        ci_str = "\n".join(
            f"{'[CINN] ' if is_cinn_kernel(k['name']) else ''}{k['name'][:100]}: {k['total_ns']}ns x{k['instances']}"
            for k in r["ci_kernels"] if not is_compute_bound(k["name"])
        )
        for c, v in enumerate([r["case"].replace("__", "/"), round(r["sp_pct"], 1), dy_str, ci_str], 1):
            cell = ws3.cell(i, c, v)
            cell.border = bd
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            cl = col[0].column_letter
            ml = max((len(str(c.value).split("\n")[0]) if c.value else 0) for c in col)
            ws.column_dimensions[cl].width = min(ml + 2, 80)

    ws1.freeze_panes = "A2"
    ws1.auto_filter.ref = ws1.dimensions
    ws3.freeze_panes = "A2"

    wb.save(OUTPUT_FILE)
    print(f"\n[DONE] 已生成 {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
